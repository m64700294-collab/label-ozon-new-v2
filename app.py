import streamlit as st
import re
import os
import csv
from io import BytesIO

from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import requests


# ============================================================
# НАСТРОЙКИ STREAMLIT
# ============================================================

st.set_page_config(
    page_title="Ozon — Этап 1 — Этикетки",
    page_icon="🖨️",
    layout="wide"
)

st.title("🖨️ Ozon — Этап 1 — Этикетки + данные API")

st.write(
    "На этом этапе программа сопоставляет обычные этикетки "
    "с выгрузкой Ozon API. Новые этикетки вида "
    "`II5010320 2549` сохраняются для второго этапа."
)


# ============================================================
# ШРИФТ
# ============================================================

@st.cache_resource
def load_font():

    font_path = "Roboto_Full_Final.ttf"

    if not os.path.exists(font_path):

        url = (
            "https://cdnjs.cloudflare.com/ajax/libs/"
            "pdfmake/0.1.66/fonts/Roboto/Roboto-Regular.ttf"
        )

        try:

            r = requests.get(
                url,
                timeout=30
            )

            r.raise_for_status()

            with open(font_path, "wb") as f:
                f.write(r.content)

        except Exception as e:

            st.error(
                f"Не удалось загрузить шрифт: {e}"
            )

            raise

    try:

        pdfmetrics.registerFont(
            TTFont(
                "OzonFont",
                font_path
            )
        )

    except Exception:
        pass

    return "OzonFont"


font_name = load_font()


# ============================================================
# НОРМАЛИЗАЦИЯ НОМЕРА ОТПРАВЛЕНИЯ
# ============================================================

def normalize_order(order):

    if order is None:
        return ""

    value = str(order).strip()

    if not value:
        return ""

    value = (
        value
        .replace("\u200b", "")
        .replace("\xa0", " ")
    )

    value = (
        value
        .replace("–", "-")
        .replace("—", "-")
        .replace("−", "-")
    )

    value = re.sub(
        r"\s*-\s*",
        "-",
        value
    )

    return value.lower()


def order_key(order):

    return normalize_order(order)


def numeric_order_key(order):

    value = normalize_order(order)

    return re.sub(
        r"\D",
        "",
        value
    )


# ============================================================
# ПОИСК НОМЕРА ОТПРАВЛЕНИЯ
# ============================================================

ORDER_PATTERN = re.compile(
    r"\d{6,15}\s*-\s*\d{2,6}\s*-\s*\d+",
    re.IGNORECASE
)


def find_orders(text):

    if not text:
        return []

    matches = ORDER_PATTERN.findall(text)

    result = []

    for value in matches:

        value = normalize_order(value)

        if value and value not in result:

            result.append(value)

    return result


# ============================================================
# ДОПОЛНИТЕЛЬНЫЙ ПОИСК
# ============================================================

def find_order_by_numeric_text(text):

    if not text:
        return []

    clean = re.sub(
        r"\D",
        "",
        text
    )

    result = []

    for length in [15, 14, 13, 12]:

        if len(clean) < length:
            continue

        for i in range(
            0,
            len(clean) - length + 1
        ):

            candidate = clean[
                i:i + length
            ]

            if candidate not in result:

                result.append(candidate)

    return result


# ============================================================
# ПОИСК КОЛОНОК
# ============================================================

def normalize_column_name(name):

    if name is None:
        return ""

    value = str(name).strip().lower()

    value = (
        value
        .replace("ё", "е")
        .replace("\xa0", " ")
    )

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value


def find_column(columns, variants):

    normalized = {}

    for column in columns:

        normalized[
            normalize_column_name(column)
        ] = column

    for variant in variants:

        v = normalize_column_name(
            variant
        )

        if v in normalized:

            return normalized[v]

    for column in columns:

        c = normalize_column_name(
            column
        )

        for variant in variants:

            v = normalize_column_name(
                variant
            )

            if v in c:

                return column

    return None


# ============================================================
# ЧТЕНИЕ XLSX
# ============================================================

def read_xlsx(file):

    try:

        import openpyxl

    except ImportError:

        raise Exception(
            "Не установлен openpyxl. "
            "Добавьте openpyxl в requirements.txt"
        )

    file.seek(0)

    workbook = openpyxl.load_workbook(
        file,
        read_only=True,
        data_only=True
    )

    sheet = workbook[
        workbook.sheetnames[0]
    ]

    rows = sheet.iter_rows(
        values_only=True
    )

    try:

        headers = next(rows)

    except StopIteration:

        return []

    headers = [
        str(h).strip()
        if h is not None
        else ""
        for h in headers
    ]

    result = []

    for row in rows:

        item = {}

        for i, header in enumerate(headers):

            if not header:
                continue

            value = (
                row[i]
                if i < len(row)
                else ""
            )

            item[header] = (
                ""
                if value is None
                else str(value).strip()
            )

        if any(
            str(v).strip()
            for v in item.values()
        ):

            result.append(item)

    workbook.close()

    return result


# ============================================================
# ЧТЕНИЕ CSV
# ============================================================

def read_csv_file(file):

    file.seek(0)

    raw = file.read()

    try:

        text = raw.decode(
            "utf-8-sig"
        )

    except UnicodeDecodeError:

        text = raw.decode(
            "cp1251",
            errors="replace"
        )

    sample = text[:5000]

    try:

        dialect = csv.Sniffer().sniff(
            sample,
            delimiters=";,|\t"
        )

        delimiter = dialect.delimiter

    except Exception:

        delimiter = ";"

    reader = csv.DictReader(
        text.splitlines(),
        delimiter=delimiter
    )

    result = []

    for row in reader:

        clean_row = {}

        for key, value in row.items():

            if key is None:
                continue

            clean_row[
                str(key).strip()
            ] = (
                ""
                if value is None
                else str(value).strip()
            )

        if any(
            str(v).strip()
            for v in clean_row.values()
        ):

            result.append(clean_row)

    return result


# ============================================================
# ЗАГРУЗКА ТАБЛИЦЫ
# ============================================================

def load_product_table(file):

    name = file.name.lower()

    if name.endswith(".xlsx"):

        rows = read_xlsx(file)

    elif name.endswith(".csv"):

        rows = read_csv_file(file)

    else:

        raise Exception(
            "Поддерживаются только XLSX и CSV."
        )

    if not rows:

        raise Exception(
            "Таблица пустая."
        )

    return rows


# ============================================================
# ИНДЕКС ОТПРАВЛЕНИЙ
# ============================================================

def build_order_index(rows):

    if not rows:

        return {}, {}, None

    columns = list(
        rows[0].keys()
    )

    order_column = find_column(
        columns,
        [
            "Номер отправления",
            "posting number",
            "posting_number",
            "Номер отправления Ozon",
            "Отправление"
        ]
    )

    article_column = find_column(
        columns,
        [
            "Артикул продавца",
            "Артикул",
            "offer_id",
            "Offer ID"
        ]
    )

    name_column = find_column(
        columns,
        [
            "Название товара",
            "Название",
            "Товар",
            "Наименование"
        ]
    )

    qty_column = find_column(
        columns,
        [
            "Количество",
            "Кол-во",
            "Кол во",
            "Qty",
            "quantity"
        ]
    )

    sku_column = find_column(
        columns,
        [
            "SKU Ozon",
            "SKU",
            "product_id"
        ]
    )

    if not order_column:

        raise Exception(
            "Не найдена колонка "
            "'Номер отправления'.\n\n"
            "Найденные колонки:\n" +
            "\n".join(
                str(x)
                for x in columns
            )
        )

    if not article_column:

        raise Exception(
            "Не найдена колонка "
            "'Артикул продавца'."
        )

    if not qty_column:

        raise Exception(
            "Не найдена колонка "
            "'Количество'."
        )

    index = {}

    numeric_index = {}

    duplicate_orders = []

    for row in rows:

        order = normalize_order(
            row.get(
                order_column,
                ""
            )
        )

        if not order:
            continue

        item = {

            "order": order,

            "article": str(
                row.get(
                    article_column,
                    ""
                )
            ).strip(),

            "name": str(
                row.get(
                    name_column,
                    ""
                )
            ).strip(),

            "qty": str(
                row.get(
                    qty_column,
                    ""
                )
            ).strip(),

            "sku": str(
                row.get(
                    sku_column,
                    ""
                )
            ).strip()
        }

        key = order_key(order)

        if key in index:

            duplicate_orders.append(
                order
            )

        index[key] = item

        numeric_key = numeric_order_key(
            order
        )

        if numeric_key:

            numeric_index[
                numeric_key
            ] = item

    return (
        index,
        numeric_index,
        duplicate_orders
    )


# ============================================================
# ПОИСК ТОВАРА
# ============================================================

def find_product(
    order,
    order_index,
    numeric_index
):

    if not order:

        return None

    info = order_index.get(
        order_key(order)
    )

    if info:

        return info

    numeric_key = numeric_order_key(
        order
    )

    if numeric_key:

        info = numeric_index.get(
            numeric_key
        )

        if info:

            return info

    return None


# ============================================================
# СОЗДАНИЕ ИНФОРМАЦИОННОЙ СТРАНИЦЫ
# ============================================================

def create_info_label(
    width,
    height,
    order_number,
    product_info,
    unresolved=False
):

    packet = BytesIO()

    c = canvas.Canvas(
        packet,
        pagesize=(
            width,
            height
        )
    )

    x_margin = 10

    # ========================================================
    # НЕ РАСПОЗНАННАЯ ЭТИКЕТКА
    # ========================================================

    if unresolved:

        c.setFont(
            font_name,
            10
        )

        c.drawString(
            x_margin,
            height - 20,
            "Заказ: НЕ РАСПОЗНАН"
        )

        c.setFont(
            font_name,
            13
        )

        c.drawString(
            x_margin,
            height - 42,
            "Арт: -"
        )

        c.setFont(
            font_name,
            24
        )

        c.drawString(
            x_margin,
            18,
            "КОЛ-ВО: ?"
        )

        c.save()

        packet.seek(0)

        return PdfReader(
            packet
        ).pages[0]

    # ========================================================
    # ОБЫЧНАЯ ЭТИКЕТКА
    # ========================================================

    c.setFont(
        font_name,
        9
    )

    c.drawString(
        x_margin,
        height - 17,
        f"Заказ: {order_number}"
    )

    c.line(
        x_margin,
        height - 20,
        width - x_margin,
        height - 20
    )

    # ========================================================
    # АРТИКУЛ
    # ========================================================

    article = str(
        product_info.get(
            "article",
            "-"
        )
    )

    c.setFont(
        font_name,
        13
    )

    if len(article) > 28:

        article = (
            article[:25] +
            "..."
        )

    c.drawString(
        x_margin,
        height - 37,
        f"Арт: {article}"
    )

    # ========================================================
    # SKU
    # ========================================================

    sku = str(
        product_info.get(
            "sku",
            ""
        )
    )

    if sku:

        c.setFont(
            font_name,
            7
        )

        c.drawString(
            x_margin,
            height - 49,
            f"SKU: {sku}"
        )

    # ========================================================
    # НАЗВАНИЕ
    # ========================================================

    name = str(
        product_info.get(
            "name",
            "Товар не найден"
        )
    )

    top_limit = height - 61

    bottom_limit = 55

    available_height = (
        top_limit -
        bottom_limit
    )

    font_size = 9

    line_height = 11

    words = name.split()

    lines = []

    current = ""

    for word in words:

        if not current:

            current = word

        elif (
            len(current)
            + len(word)
            + 1
            <= 38
        ):

            current += " " + word

        else:

            lines.append(
                current
            )

            current = word

    if current:

        lines.append(
            current
        )

    while (
        len(lines) * line_height
        > available_height
        and font_size > 5.5
    ):

        font_size -= 0.5

        line_height -= 0.5

        max_chars = int(
            38 *
            (9 / font_size)
        )

        lines = []

        current = ""

        for word in name.split():

            if not current:

                current = word

            elif (
                len(current)
                + len(word)
                + 1
                <= max_chars
            ):

                current += " " + word

            else:

                lines.append(
                    current
                )

                current = word

        if current:

            lines.append(
                current
            )

    c.setFont(
        font_name,
        font_size
    )

    y = top_limit

    for line in lines:

        if y <= bottom_limit:

            break

        c.drawString(
            x_margin,
            y,
            line
        )

        y -= line_height

    # ========================================================
    # КОЛИЧЕСТВО
    # ========================================================

    qty = str(
        product_info.get(
            "qty",
            "?"
        )
    )

    c.setFont(
        font_name,
        24
    )

    c.drawString(
        x_margin,
        15,
        f"КОЛ-ВО: {qty}"
    )

    c.save()

    packet.seek(0)

    return PdfReader(
        packet
    ).pages[0]


# ============================================================
# ИНТЕРФЕЙС
# ============================================================

st.divider()

col1, col2 = st.columns(2)

with col1:

    labels_file = st.file_uploader(
        "1️⃣ PDF этикеток Ozon",
        type=["pdf"]
    )

with col2:

    table_file = st.file_uploader(
        "2️⃣ Выгрузка Ozon API",
        type=["xlsx", "csv"]
    )


# ============================================================
# ЧТЕНИЕ ТАБЛИЦЫ
# ============================================================

table_rows = None

if table_file:

    try:

        table_rows = load_product_table(
            table_file
        )

        st.success(
            f"✅ Таблица прочитана: "
            f"{len(table_rows)} строк"
        )

        with st.expander(
            "🔎 Предпросмотр"
        ):

            st.dataframe(
                table_rows[:20],
                use_container_width=True
            )

    except Exception as e:

        st.error(
            f"Ошибка чтения таблицы: {e}"
        )

        table_rows = None


# ============================================================
# ОБРАБОТКА
# ============================================================

if (
    labels_file
    and table_file
    and table_rows
):

    if st.button(
        "🚀 ЭТАП 1 — СОЗДАТЬ PDF",
        type="primary",
        use_container_width=True
    ):

        try:

            (
                order_index,
                numeric_index,
                duplicate_orders
            ) = build_order_index(
                table_rows
            )

            labels_file.seek(0)

            reader = PdfReader(
                labels_file
            )

            writer = PdfWriter()

            total_labels = len(
                reader.pages
            )

            success_count = 0
            unresolved_count = 0

            unresolved_pages = []

            progress = st.progress(0)

            status = st.empty()

            # =================================================
            # СТРАНИЦЫ
            # =================================================

            for i, page in enumerate(
                reader.pages
            ):

                status.text(
                    f"🏷 Этикетка "
                    f"{i + 1} / "
                    f"{total_labels}"
                )

                # -------------------------------------------------
                # САМА Ozon-ЭТИКЕТКА
                # -------------------------------------------------

                writer.add_page(
                    page
                )

                try:

                    text = (
                        page.extract_text()
                        or ""
                    )

                except Exception:

                    text = ""

                # -------------------------------------------------
                # ИЩЕМ ПОЛНЫЙ НОМЕР
                # -------------------------------------------------

                orders = find_orders(
                    text
                )

                full_order = (
                    orders[0]
                    if orders
                    else ""
                )

                # -------------------------------------------------
                # РЕЗЕРВНЫЙ ПОИСК
                # -------------------------------------------------

                if not full_order:

                    numeric_candidates = (
                        find_order_by_numeric_text(
                            text
                        )
                    )

                    for candidate in (
                        numeric_candidates
                    ):

                        if candidate in numeric_index:

                            info = (
                                numeric_index[
                                    candidate
                                ]
                            )

                            full_order = (
                                info["order"]
                            )

                            break

                # -------------------------------------------------
                # ИЩЕМ ТОВАР
                # -------------------------------------------------

                info = None

                if full_order:

                    info = find_product(
                        full_order,
                        order_index,
                        numeric_index
                    )

                # =================================================
                # ОБЫЧНАЯ ЭТИКЕТКА
                # =================================================

                if info:

                    success_count += 1

                    info_page = (
                        create_info_label(
                            width=float(
                                page.mediabox.width
                            ),
                            height=float(
                                page.mediabox.height
                            ),
                            order_number=info["order"],
                            product_info=info,
                            unresolved=False
                        )
                    )

                # =================================================
                # НОВАЯ ЭТИКЕТКА
                # =================================================

                else:

                    unresolved_count += 1

                    unresolved_pages.append(
                        {
                            "page": i + 1,
                            "text": text[:1000]
                        }
                    )

                    info_page = (
                        create_info_label(
                            width=float(
                                page.mediabox.width
                            ),
                            height=float(
                                page.mediabox.height
                            ),
                            order_number="",
                            product_info={},
                            unresolved=True
                        )
                    )

                # -------------------------------------------------
                # ВТОРАЯ СТРАНИЦА
                # -------------------------------------------------

                writer.add_page(
                    info_page
                )

                progress.progress(
                    (i + 1) /
                    total_labels
                )

            status.text(
                "✅ Этап 1 завершён"
            )

            # ====================================================
            # СТАТИСТИКА
            # ====================================================

            st.divider()

            c1, c2, c3 = st.columns(3)

            c1.metric(
                "Всего этикеток",
                total_labels
            )

            c2.metric(
                "Найдено в API",
                success_count
            )

            c3.metric(
                "На второй этап",
                unresolved_count
            )

            if unresolved_count:

                st.warning(
                    f"⚠️ {unresolved_count} "
                    "этикеток не распознаны. "
                    "Они сохранены в PDF с пометкой "
                    "'НЕ РАСПОЗНАН' и будут обработаны "
                    "на втором этапе по PDF листа подбора."
                )

            else:

                st.success(
                    "🎉 Все этикетки распознаны "
                    "на первом этапе."
                )

            # ====================================================
            # PDF
            # ====================================================

            output = BytesIO()

            writer.write(
                output
            )

            output.seek(0)

            st.download_button(
                "📥 Скачать PDF для Этапа 2",
                output,
                "Ozon_Etap_1.pdf",
                "application/pdf",
                type="primary",
                use_container_width=True
            )

            # ====================================================
            # ДИАГНОСТИКА
            # ====================================================

            if unresolved_pages:

                with st.expander(
                    "🔎 Что отправлено на Этап 2"
                ):

                    for item in unresolved_pages:

                        st.write(
                            f"Страница PDF: "
                            f"{item['page']}"
                        )

                        if item["text"]:

                            st.code(
                                item["text"]
                            )

                        st.divider()

        except Exception as e:

            st.error(
                "❌ Ошибка обработки:\n\n"
                + str(e)
            )

            with st.expander(
                "Техническая информация"
            ):

                st.exception(e)


# ============================================================
# ПОДСКАЗКА
# ============================================================

if not labels_file or not table_file:

    st.info(
        """
        ### Этап 1

        Загрузите:

        **1.** PDF с этикетками Ozon.

        **2.** CSV/XLSX выгрузку Ozon API.

        Обычные этикетки вида:

        `0126236473-0731-1`

        будут сразу сопоставлены.

        Этикетки нового формата:

        `II5010320 2549`

        не удаляются и не пытаются сопоставиться
        по неправильному номеру.

        Они переходят на **Этап 2**, где номер
        `2549` будет найден в PDF листа подбора
        и связан с номером отправления Ozon.
        """
    )
